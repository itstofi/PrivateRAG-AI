import mimetypes
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import fitz
from docx import Document as DocxDocument

from app.core.exceptions import ValidationError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
ALLOWED_MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
    },
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
}


@dataclass(frozen=True)
class ExtractedPage:
    text: str
    page_number: int | None
    section: str | None = None


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def validate_upload(filename: str, content: bytes, max_bytes: int, content_type: str | None) -> str:
    if (
        not filename
        or Path(filename).name != filename
        or "/" in filename
        or "\\" in filename
        or CONTROL_FILENAME.search(filename)
    ):
        raise ValidationError("The uploaded filename is invalid.")
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValidationError("Supported file types are PDF, DOCX, TXT, and Markdown.")
    if not content:
        raise ValidationError("The uploaded file is empty.")
    if len(content) > max_bytes:
        raise ValidationError(
            f"The uploaded file exceeds the {max_bytes // 1024 // 1024} MB limit."
        )
    declared = (content_type or "").split(";")[0].strip().lower()
    if declared and declared not in ALLOWED_MIME_TYPES[extension]:
        raise ValidationError("The file's MIME type does not match its extension.")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise ValidationError("The PDF signature is invalid.")
    if extension == ".docx" and not content.startswith(b"PK"):
        raise ValidationError("The DOCX archive signature is invalid.")
    if extension == ".docx":
        validate_docx_archive(content, max_bytes)
    if extension in {".txt", ".md"}:
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationError("Text and Markdown files must use UTF-8 encoding.") from exc
    return extension


CONTROL_FILENAME = re.compile(r"[\x00-\x1f\x7f]")


def validate_docx_archive(content: bytes, compressed_limit: int) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ValidationError("The DOCX archive is missing required document parts.")
            total_uncompressed = sum(item.file_size for item in archive.infolist())
            if len(names) > 10_000 or total_uncompressed > compressed_limit * 20:
                raise ValidationError("The DOCX archive expands beyond the safe processing limit.")
    except BadZipFile as exc:
        raise ValidationError("The DOCX archive is malformed.") from exc


def load_pdf(path: Path, max_pages: int = 1000, max_chars: int = 2_000_000) -> list[ExtractedPage]:
    try:
        with fitz.open(path) as document:
            if document.page_count > max_pages:
                raise ValidationError(f"The PDF exceeds the {max_pages}-page safety limit.")
            pages: list[ExtractedPage] = []
            extracted_chars = 0
            for page in document:
                text = clean_text(page.get_text("text"))
                extracted_chars += len(text)
                if extracted_chars > max_chars:
                    raise ValidationError("The document contains too much extracted text.")
                pages.append(ExtractedPage(text, page.number + 1))
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("The PDF could not be read.", detail=str(exc)) from exc
    return [page for page in pages if page.text]


def load_docx(path: Path, max_chars: int = 2_000_000) -> list[ExtractedPage]:
    try:
        document = DocxDocument(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = clean_text("\n".join(blocks))
        if len(text) > max_chars:
            raise ValidationError("The document contains too much extracted text.")
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("The DOCX file could not be read.", detail=str(exc)) from exc
    return [ExtractedPage(text, None, "Document body")] if text else []


def load_text(path: Path, max_chars: int = 2_000_000) -> list[ExtractedPage]:
    try:
        text = clean_text(path.read_text(encoding="utf-8-sig"))
        if len(text) > max_chars:
            raise ValidationError("The document contains too much extracted text.")
    except (UnicodeDecodeError, OSError) as exc:
        raise ValidationError("The text file must use UTF-8 encoding.", detail=str(exc)) from exc
    return [ExtractedPage(text, None, "Document body")] if text else []


def load_document(
    path: Path, max_pages: int = 1000, max_chars: int = 2_000_000
) -> list[ExtractedPage]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        pages = load_pdf(path, max_pages, max_chars)
    elif extension == ".docx":
        pages = load_docx(path, max_chars)
    elif extension in {".txt", ".md"}:
        pages = load_text(path, max_chars)
    else:
        guessed = mimetypes.guess_type(path.name)[0]
        raise ValidationError(f"Unsupported document format ({guessed or 'unknown'}).")
    if not pages:
        raise ValidationError("No readable text was found in the document.")
    return pages
