from pathlib import Path

import fitz
import pytest
from docx import Document

from app.core.exceptions import ValidationError
from app.ingestion.chunking import chunk_pages, split_text
from app.ingestion.loaders import ExtractedPage, load_docx, load_pdf, validate_upload


def test_upload_validation() -> None:
    assert validate_upload("notes.md", b"# Local notes", 1000, "text/markdown") == ".md"
    with pytest.raises(ValidationError):
        validate_upload("../notes.md", b"text", 1000, "text/markdown")
    with pytest.raises(ValidationError):
        validate_upload("report.pdf", b"not a pdf", 1000, "application/pdf")
    with pytest.raises(ValidationError):
        validate_upload("large.txt", b"x" * 11, 10, "text/plain")
    with pytest.raises(ValidationError):
        validate_upload("..\\notes.txt", b"text", 1000, "text/plain")
    with pytest.raises(ValidationError):
        validate_upload("binary.txt", b"\xff\xfe", 1000, "text/plain")


def test_pdf_text_extraction_preserves_pages(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    document = fitz.open()
    for text in ("First page policy", "Second page allowance"):
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    pages = load_pdf(path)
    assert [page.page_number for page in pages] == [1, 2]
    assert "allowance" in pages[1].text


def test_docx_text_extraction_includes_table(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Local policy")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Allowance"
    table.rows[0].cells[1].text = "Two days"
    document.save(path)
    pages = load_docx(path)
    assert "Local policy" in pages[0].text
    assert "Allowance | Two days" in pages[0].text


def test_chunking_overlaps_and_numbers() -> None:
    text = " ".join(f"word{index}" for index in range(80))
    parts = split_text(text, 100, 20)
    assert len(parts) > 2
    chunks = chunk_pages([ExtractedPage(text, 7)], 100, 20)
    assert chunks[0].page_number == 7
    assert [chunk.chunk_number for chunk in chunks] == list(range(1, len(chunks) + 1))
    assert all(chunk.text for chunk in chunks)
